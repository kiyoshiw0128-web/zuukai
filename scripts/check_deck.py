"""できた .pptx を検査する。zuukai_core で作っていない資料にも使える。

    python check_deck.py 資料.pptx [--min-narration 40] [--script 台本.md]

止める（error）: ナレーションがないスライド、18pt未満の文字、スライドの外にはみ出した図形
見直す（warn） : 文字が箱に収まらなさそう、1枚の文字が多い、ナレーションが長すぎ・短すぎ（100字未満、2分半超）、
                 ナレーションが話し言葉（です・ます）になっていない
--script を付けると、ナレーションだけを並べた読み上げ台本を書き出す。
終了コード: error があれば 1、なければ 0。
"""
from __future__ import annotations

import argparse
import math
import re
import sys
import unicodedata
from dataclasses import dataclass

from pptx import Presentation
from pptx.util import Emu

MIN_PT = 18
MAX_SLIDE_CHARS = 120       # スライド上の文字の目安（ノートは数えない）
SPEAK_CHARS_PER_MIN = 300   # 日本語の話す速さの目安
EMU_PER_IN = 914400


@dataclass(frozen=True)
class Finding:
    level: str
    slide: int
    message: str


def char_width(ch: str) -> float:
    """全角を1、半角を0.55として数える。"""
    return 1.0 if unicodedata.east_asian_width(ch) in "WFA" else 0.55


def text_frames(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # グループ
            yield from text_frames(sh.shapes)
        if sh.has_text_frame:
            yield sh, sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield None, cell.text_frame


def estimate_overflow(sh, tf) -> bool:
    """箱の幅と高さから、必要な行数が入るかを大まかに見積もる。"""
    if sh is None or not sh.width or not sh.height or not tf.text.strip():
        return False
    width_pt = Emu(sh.width).inches * 72 - 8
    height_pt = Emu(sh.height).inches * 72
    need = 0.0
    for p in tf.paragraphs:
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=18)
        w = sum(char_width(c) for c in "".join(r.text for r in p.runs)) * size
        lines = max(1, math.ceil(w / max(width_pt, 1))) if tf.word_wrap is not False else 1
        need += lines * size * 1.2
    return need > height_pt * 1.15


def check_slide(n: int, slide, prs, min_narration: int) -> list[Finding]:
    out: list[Finding] = []
    sw, shh = prs.slide_width, prs.slide_height
    chars = 0
    for sh, tf in text_frames(slide.shapes):
        for p in tf.paragraphs:
            for r in p.runs:
                if r.font.size and r.font.size.pt < MIN_PT and r.text.strip():
                    out.append(Finding("error", n, f"{r.font.size.pt:g}pt の文字: 「{r.text[:20]}」"))
            chars += len("".join(r.text for r in p.runs).strip())
        if estimate_overflow(sh, tf):
            out.append(Finding("warn", n, f"箱に収まらなさそう: 「{tf.text[:24]}」"))
    for sh in slide.shapes:
        if sh.left is None or sh.width is None:
            continue
        if sh.left < -EMU_PER_IN * 0.01 or sh.top < -EMU_PER_IN * 0.01 or \
                sh.left + sh.width > sw * 1.001 or sh.top + sh.height > shh * 1.001:
            out.append(Finding("error", n, f"スライドの外にはみ出した図形: {sh.name}"))
    if chars > MAX_SLIDE_CHARS:
        out.append(Finding("warn", n, f"スライド上の文字が{chars}字。図に置き換えるか、ノートへ移す"))

    notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
    if len(notes) < min_narration:
        out.append(Finding("error", n, f"ナレーションがない・短い（{len(notes)}字）"))
    else:
        sec = len(notes) / SPEAK_CHARS_PER_MIN * 60
        if len(notes) < 100:
            out.append(Finding("warn", n, f"ナレーションが{len(notes)}字。100字（約20秒）を目安に、"
                                          "どこを見るか・何が言えるか・つなぎを足す"))
        if sec > 150:
            out.append(Finding("warn", n, f"ナレーションが約{sec:.0f}秒。2分半を超えるので分けるか削る"))
        if not re.search(r"(です|ます|でした|ました|ません|ください)[。！？]?", notes):
            out.append(Finding("warn", n, "ナレーションが話し言葉（です・ます）になっていない"))
    return out


def write_script(prs, path: str) -> None:
    lines = ["# 読み上げ台本", ""]
    total = 0
    for n, s in enumerate(prs.slides, 1):
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        total += len(notes)
        lines += [f"## {n}枚目（約{len(notes) / SPEAK_CHARS_PER_MIN * 60:.0f}秒）", "", notes, ""]
    lines.insert(1, f"\n全体で約{total / SPEAK_CHARS_PER_MIN:.1f}分\n")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("pptx")
    ap.add_argument("--min-narration", type=int, default=40)
    ap.add_argument("--script", help="ナレーションを並べた台本 .md の書き出し先")
    args = ap.parse_args(argv)
    sys.stdout.reconfigure(encoding="utf-8")

    prs = Presentation(args.pptx)
    findings = [f for n, s in enumerate(prs.slides, 1)
                for f in check_slide(n, s, prs, args.min_narration)]
    errors = [f for f in findings if f.level == "error"]
    for f in findings:
        print(f"[{f.level.upper():5}] {f.slide:2}枚目: {f.message}")
    print(f"\n{len(prs.slides)}枚を検査: error {len(errors)}件、warn {len(findings) - len(errors)}件")
    if args.script:
        write_script(prs, args.script)
        print(f"台本を書き出しました: {args.script}")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
