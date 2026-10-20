"""PowerPoint がなくても配置を目で確かめるための、簡易プレビュー（HTML）を書き出す。

    python preview_html.py 資料.pptx 出力.html

四角・角丸・丸・線・文字だけを描く簡易版です。最終の見た目は PowerPoint で確かめてください。
"""
from __future__ import annotations

import html
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

PX_PER_IN = 72  # 1インチ＝72px（文字の pt とそろう）


def _color(fill) -> str | None:
    try:
        if fill.type == 1:
            return f"#{fill.fore_color.rgb}"
    except (AttributeError, TypeError, KeyError):
        return None
    return None


def _line_color(sh) -> str | None:
    try:
        if sh.line.fill.type == 1:
            return f"#{sh.line.color.rgb}"
    except (AttributeError, TypeError, KeyError):
        return None
    return None


def _px(v) -> float:
    return Emu(v).inches * PX_PER_IN


def shape_html(sh) -> str:
    x, y, w, h = _px(sh.left), _px(sh.top), _px(sh.width), _px(sh.height)
    if sh.shape_type == MSO_SHAPE_TYPE.LINE or sh.shape_type == MSO_SHAPE_TYPE.FREEFORM:
        bx, by, ex, ey = _px(sh.begin_x), _px(sh.begin_y), _px(sh.end_x), _px(sh.end_y)
        c = _line_color(sh) or "#888"
        return (f'<svg class="ln"><line x1="{bx}" y1="{by}" x2="{ex}" y2="{ey}" stroke="{c}" '
                f'stroke-width="2.5" marker-end="url(#ah)"/></svg>') if hasattr(sh, "begin_x") else ""
    style = [f"left:{x}px;top:{y}px;width:{w}px;height:{h}px"]
    fill = _color(sh.fill) if hasattr(sh, "fill") else None
    if fill:
        style.append(f"background:{fill}")
    lc = _line_color(sh) if hasattr(sh, "line") else None
    if lc:
        style.append(f"border:1.5px solid {lc}")
    kind = str(getattr(sh, "auto_shape_type", "")) if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE else ""
    if "OVAL" in kind:
        style.append("border-radius:50%")
    elif "ROUND" in kind:
        style.append("border-radius:8px")
    inner = ""
    if sh.has_text_frame and sh.text_frame.text.strip():
        anchor = {"MIDDLE (3)": "center", "BOTTOM (4)": "flex-end"}.get(
            str(sh.text_frame.vertical_anchor), "flex-start")
        paras = []
        for p in sh.text_frame.paragraphs:
            runs = p.runs
            if not runs:
                continue
            f = runs[0].font
            size = f.size.pt if f.size else 18
            col = f"#{f.color.rgb}" if f.color and f.color.type else "#262626"
            align = {"CENTER (2)": "center", "RIGHT (3)": "right"}.get(str(p.alignment), "left")
            paras.append(f'<p style="font-size:{size}px;color:{col};text-align:{align};'
                         f'font-weight:{700 if f.bold else 400}">'
                         f'{html.escape("".join(r.text for r in runs))}</p>')
        inner = f'<div class="tx" style="justify-content:{anchor}">{"".join(paras)}</div>'
    return f'<div class="sh" style="{";".join(style)}">{inner}</div>'


def main(src: str, dst: str) -> None:
    prs = Presentation(src)
    W, H = _px(prs.slide_width), _px(prs.slide_height)
    out = [f"""<!doctype html><meta charset="utf-8"><title>preview</title><style>
body{{background:#888;font-family:'Noto Sans JP','Meiryo',sans-serif;margin:0;padding:16px}}
.sl{{position:relative;width:{W}px;height:{H}px;background:#fff;margin:0 auto 12px;overflow:hidden}}
.sh{{position:absolute;box-sizing:border-box}} .ln{{position:absolute;left:0;top:0;width:{W}px;height:{H}px;overflow:visible}}
.tx{{display:flex;flex-direction:column;height:100%;padding:2px}} .tx p{{margin:0;line-height:1.2}}
.nt{{width:{W}px;margin:0 auto 28px;background:#fffbe6;padding:8px;font-size:14px;white-space:pre-wrap}}
</style><svg width="0" height="0"><defs><marker id="ah" markerWidth="8" markerHeight="8" refX="6" refY="3" orient="auto">
<path d="M0,0 L6,3 L0,6 z" fill="#888"/></marker></defs></svg>"""]
    for n, s in enumerate(prs.slides, 1):
        body = "".join(shape_html(sh) for sh in s.shapes)
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else "（ナレーションなし）"
        out.append(f'<div class="sl">{body}</div><div class="nt"><b>{n}枚目 ナレーション</b>\n'
                   f"{html.escape(notes)}</div>")
    with open(dst, "w", encoding="utf-8") as f:
        f.write("\n".join(out))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("使い方: python preview_html.py 資料.pptx 出力.html")
    main(sys.argv[1], sys.argv[2])
