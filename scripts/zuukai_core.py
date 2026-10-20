"""zuukai の土台: デッキ、色と文字の決まり、基本の図形、アイコン。

使い方の例は examples/build_example.py にあります。
座標と大きさの単位はすべてインチ（16:9、13.333 x 7.5）です。
"""
from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---- 色と文字の決まり（design tokens） ----------------------------------

PALETTE = {
    "navy": "1F3A5F",    # 人・結論・見出し
    "accent": "C8743A",  # AI・注目させたい所（1枚に1か所まで）
    "paper": "F3F1EC",   # 図を載せる下地
    "ink": "262626",     # 本文
    "gray": "6B6B6B",    # 補足・出典・「いいえ」
    "line": "C8C8C8",    # 罫線・使わない物
    "mist": "E6EBF1",    # 選ばれた物の薄い塗り
    "brown": "906146",   # 段階の2番目
    "plum": "574D53",    # 段階の3番目
    "white": "FFFFFF",
}
FONT = "Noto Sans JP"
MIN_PT = 18          # 投影で読める最小の文字。これ未満は検査で止める
SLIDE_W, SLIDE_H = 13.333, 7.5

# 標準の枠（見出しの下から提案の帯の上まで）
CONTENT = (0.7, 1.5, 11.93, 4.1)
TITLE_BOX = (0.7, 0.5, 11.93, 0.9)
BAND_Y = 5.8


def rgb(name_or_hex: str) -> RGBColor:
    return RGBColor.from_string(PALETTE.get(name_or_hex, name_or_hex))


@dataclass
class Box:
    x: float
    y: float
    w: float
    h: float

    @classmethod
    def of(cls, b) -> "Box":
        return b if isinstance(b, Box) else cls(*b)

    @property
    def cx(self) -> float:
        return self.x + self.w / 2

    @property
    def cy(self) -> float:
        return self.y + self.h / 2

    def split_cols(self, n: int, gap: float = 0.3) -> list["Box"]:
        w = (self.w - gap * (n - 1)) / n
        return [Box(self.x + i * (w + gap), self.y, w, self.h) for i in range(n)]

    def split_rows(self, n: int, gap: float = 0.2) -> list["Box"]:
        h = (self.h - gap * (n - 1)) / n
        return [Box(self.x, self.y + i * (h + gap), self.w, h) for i in range(n)]

    def inset(self, d: float) -> "Box":
        return Box(self.x + d, self.y + d, self.w - 2 * d, self.h - 2 * d)


# ---- デッキ -------------------------------------------------------------

@dataclass
class Deck:
    """ナレーション（ノート）なしのスライドは作れないデッキ。"""

    font: str = FONT
    min_narration_chars: int = 40
    prs: Presentation = field(default_factory=Presentation)

    def __post_init__(self):
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self._blank = self.prs.slide_layouts[6]

    def _new(self, narration: str):
        text = (narration or "").strip()
        if len(text) < self.min_narration_chars:
            raise ValueError(
                f"ナレーションが短すぎます（{len(text)}字）。各スライドに"
                f"{self.min_narration_chars}字以上の話し言葉の台本を入れてください。"
            )
        s = self.prs.slides.add_slide(self._blank)
        s.notes_slide.notes_text_frame.text = text
        return s

    def cover(self, title: str, subtitle: str, narration: str,
              footnote: str = "", art=None) -> "Slide":
        """表紙。art は (Slide, Box) を受け取って絵を描く関数。"""
        sl = Slide(self, self._new(narration))
        sl.rect(Box(0, 0, SLIDE_W, SLIDE_H), "paper")
        sl.rect(Box(0, 0, 0.35, SLIDE_H), "navy")
        sl.text(Box(1.0, 1.9, 7.6, 1.2), title, 48, "navy", bold=True)
        sl.text(Box(1.0, 3.2, 7.6, 0.8), subtitle, 28, "ink")
        sl.rect(Box(1.0, 4.35, 6.6, 0.14), "accent", rounded=True)
        if footnote:
            sl.text(Box(1.0, 5.6, 7.4, 0.5), footnote, 20, "gray")
        if art:
            art(sl, Box(8.3, 2.2, 4.3, 2.8))
        return sl

    def slide(self, title: str, narration: str) -> "Slide":
        """本文スライド。title は「言い切りの1文」にする（例: 急ぐと正答が下がる）。"""
        sl = Slide(self, self._new(narration))
        sl.rect(Box(0, 0, 0.18, SLIDE_H), "navy")
        sl.text(Box.of(TITLE_BOX), title, 36, "navy", bold=True, anchor="middle")
        return sl

    def save(self, path: str) -> None:
        for i, s in enumerate(self.prs.slides, 1):
            if not s.has_notes_slide or not s.notes_slide.notes_text_frame.text.strip():
                raise ValueError(f"{i}枚目にナレーションがありません")
        self.prs.save(path)


# ---- スライドと基本の図形 -----------------------------------------------

class Slide:
    def __init__(self, deck: Deck, s):
        self.deck, self.s = deck, s

    # 図形
    def rect(self, b, fill="paper", rounded=False, line=None, shape=None):
        b = Box.of(b)
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
        sh = self.s.shapes.add_shape(kind, Inches(b.x), Inches(b.y), Inches(b.w), Inches(b.h))
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = min(0.5, 0.12 / max(min(b.w, b.h), 0.01) * 2)
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(1.5)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def oval(self, cx, cy, d, fill="navy", line=None):
        sh = self.rect(Box(cx - d / 2, cy - d / 2, d, d), fill, shape=MSO_SHAPE.OVAL, line=line)
        return sh

    def text(self, b, text, size=20, color="ink", bold=False, align="left",
             anchor="top", shape=None, fill=None):
        """文字。size は MIN_PT 未満にしない。shape を渡すとその図形の中に書く。"""
        if size < MIN_PT:
            raise ValueError(f"文字が小さすぎます（{size}pt）。{MIN_PT}pt以上にしてください: {text[:20]}")
        b = Box.of(b)
        if shape is None:
            shape = self.s.shapes.add_textbox(Inches(b.x), Inches(b.y), Inches(b.w), Inches(b.h))
            if fill:
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(fill)
        tf = shape.text_frame
        tf.word_wrap = True
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", Inches(0.04))
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                              "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        for i, line in enumerate(str(text).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                           "right": PP_ALIGN.RIGHT}[align]
            r = p.add_run()
            r.text = line
            f = r.font
            f.size, f.bold, f.name = Pt(size), bold, self.deck.font
            f.color.rgb = rgb(color)
            _set_ea_font(r, self.deck.font)
        return shape

    def label(self, b, text, fill="navy", color="white", size=20, bold=True):
        """塗りつぶした角丸に中央ぞろえの文字（札・ボタン・箱）。

        いちばん長い行が1行に収まるよう、MIN_PT まで文字を小さくする。
        それでも収まらないときは折り返すので、言葉を短くする。
        """
        b = Box.of(b)
        sh = self.rect(b, fill, rounded=True)
        return self.text(b, text, fit_size(text, b.w, size), color, bold, "center", "middle",
                         shape=sh)

    def arrow(self, x1, y1, x2, y2, color="gray", width=2.25, head=True):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                        Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(width)
        if head:
            ln = c.line._get_or_add_ln()
            tail = etree.SubElement(ln, qn("a:tailEnd"))
            tail.set("type", "triangle")
            tail.set("w", "med")
            tail.set("h", "med")
        return c

    # 決まった部品
    def tag(self, x, y, kind="根拠"):
        """論文や数字の結果には「根拠」、自分たちの考えには「提案」の札を付ける。"""
        fill, color = ("gray", "white") if kind == "根拠" else ("white", "navy")
        return self.label(Box(x, y, 0.85, 0.42), kind, fill, color, 18)

    def panel(self, b=CONTENT, caption=None, tag=None):
        """図を載せる下地。caption は左上の小見出し（何の図か）。"""
        b = Box.of(b)
        self.rect(b, "paper", rounded=True)
        x = b.x + 0.25
        if tag:
            self.tag(x, b.y + 0.2, tag)
            x += 1.0
        if caption:
            self.text(Box(x, b.y + 0.18, b.x + b.w - x - 0.25, 0.45), caption, 20, "gray", True)
        return Box(b.x + 0.25, b.y + 0.75, b.w - 0.5, b.h - 0.95)

    def band(self, message, tag="提案"):
        """下の帯: このスライドから「だから何をするか」を1文で。"""
        self.rect(Box(0.7, BAND_Y, 11.93, 0.85), "navy", rounded=True)
        x = 1.0
        if tag:
            self.tag(x, BAND_Y + 0.21, tag)
            x = 2.15
        self.text(Box(x, BAND_Y, 12.33 - x, 0.85), message, 24, "white", True, anchor="middle")

    def source(self, text):
        self.text(Box(0.7, 6.95, 11.93, 0.4), text, 18, "gray")


def text_units(line: str) -> float:
    """全角を1、半角を0.55として、1行の幅を文字数で数える。"""
    return sum(1.0 if unicodedata.east_asian_width(c) in "WFA" else 0.55 for c in line)


def fit_size(text: str, width_in: float, size: int) -> int:
    """width_in の幅に最長の行が収まる文字の大きさ（MIN_PT より下げない）。"""
    longest = max((text_units(t) for t in str(text).split("\n")), default=1) or 1
    room_pt = width_in * 72 - 10
    return max(MIN_PT, min(size, int(room_pt / longest)))


def _set_ea_font(run, font):
    """日本語の文字にも同じフォントを当てる（latin だけだと和文が置き換わる）。"""
    rpr = run._r.get_or_add_rPr()
    for tag_name in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag_name))
        if el is None:
            el = etree.SubElement(rpr, qn(tag_name))
        el.set("typeface", font)


# ---- アイコン（図形だけで描く。画像ファイルを使わない） --------------------

def icon_ai(sl: Slide, cx, cy, size=1.0, color="accent"):
    """AI（ロボットの顔）。"""
    u = size
    sl.oval(cx, cy - 0.5 * u, 0.2 * u, color)
    sl.rect(Box(cx - 0.035 * u, cy - 0.42 * u, 0.07 * u, 0.2 * u), color)
    sl.rect(Box(cx - 0.62 * u, cy - 0.25 * u, 1.24 * u, 0.98 * u), color, rounded=True)
    for side in (-1, 1):
        sl.rect(Box(cx + side * 0.68 * u - 0.06 * u, cy + 0.05 * u, 0.12 * u, 0.35 * u), color)
        sl.oval(cx + side * 0.3 * u, cy + 0.15 * u, 0.24 * u, "white")
    sl.rect(Box(cx - 0.3 * u, cy + 0.48 * u, 0.6 * u, 0.09 * u), "white")


def icon_person(sl: Slide, cx, cy, size=1.0, color="navy"):
    """人（頭と肩）。"""
    u = size
    sl.oval(cx, cy - 0.3 * u, 0.62 * u, color)
    sl.rect(Box(cx - 0.6 * u, cy + 0.08 * u, 1.2 * u, 0.7 * u), color,
            shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)


def icon_doc(sl: Slide, cx, cy, size=1.0, color="accent"):
    """書類（下書き・報告）。"""
    u = size
    sl.rect(Box(cx - 0.45 * u, cy - 0.6 * u, 0.9 * u, 1.2 * u), "white",
            shape=MSO_SHAPE.FOLDED_CORNER, line="line")
    for i, w in enumerate((0.62, 0.62, 0.62, 0.4)):
        sl.rect(Box(cx - 0.31 * u, cy - 0.38 * u + i * 0.2 * u, w * u, 0.06 * u), color)


def icon_clock(sl: Slide, cx, cy, size=1.0, color="navy"):
    """時計（時間・締切）。"""
    u = size
    sl.oval(cx, cy, 1.3 * u, color)
    sl.oval(cx, cy, 1.02 * u, "white")
    sl.arrow(cx, cy, cx, cy - 0.38 * u, color, 3, head=False)
    sl.arrow(cx, cy, cx + 0.28 * u, cy, color, 3, head=False)
    sl.oval(cx, cy, 0.1 * u, color)


def icon_check(sl: Slide, cx, cy, size=1.0, ok=True):
    """丸に ○（よい）か ×（だめ）。"""
    color = "navy" if ok else "accent"
    sl.oval(cx, cy, 0.8 * size, color)
    sl.text(Box(cx - 0.4 * size, cy - 0.4 * size, 0.8 * size, 0.8 * size),
            "○" if ok else "×", max(MIN_PT, int(32 * size)), "white", True, "center", "middle")


ICONS = {"ai": icon_ai, "person": icon_person, "doc": icon_doc, "clock": icon_clock}
