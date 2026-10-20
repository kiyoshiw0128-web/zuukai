"""zuukai のテスト。  python -m pytest tests"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "examples"))

import check_deck  # noqa: E402
from zuukai_core import CONTENT, Box, Deck, fit_size  # noqa: E402
from zuukai_parts import bars, decision_flow  # noqa: E402

NARRATION = "左の棒を見てください。議題を送らなかった会議は平均52分でした。右は31分です。だから議題は前日に送ります。"


def test_slide_without_narration_is_rejected():
    with pytest.raises(ValueError, match="ナレーション"):
        Deck().slide("見出し", "")


def test_small_text_is_rejected():
    sl = Deck().slide("見出し", NARRATION)
    with pytest.raises(ValueError, match="18pt"):
        sl.text(Box(1, 1, 3, 1), "小さい", 12)


def test_fit_size_shrinks_long_labels_but_not_below_min():
    assert fit_size("短い", 3.0, 22) == 22
    assert fit_size("決めることから話す", 2.5, 22) < 22
    assert fit_size("とても長いラベルをこの狭い箱に入れようとする", 1.0, 22) == 18


def test_decision_flow_needs_one_more_outcome():
    sl = Deck().slide("見出し", NARRATION)
    with pytest.raises(ValueError):
        decision_flow(sl, CONTENT, ["問い？"], ["行き先"])


def test_built_deck_passes_check(tmp_path):
    d = Deck()
    sl = d.slide("議題を先に送ると、会議は短くなる", NARRATION)
    bars(sl, sl.panel(CONTENT, "会議の平均時間（分）", tag="根拠"),
         [("議題なし", 52, "gray"), ("議題を送る", 31, "accent")], unit="分")
    sl.band("だから、議題は前日までに送る")
    out = tmp_path / "ok.pptx"
    d.save(str(out))
    assert check_deck.main([str(out)]) == 0


def test_check_flags_missing_narration_and_small_text(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "小さい文字"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    out = tmp_path / "bad.pptx"
    prs.save(str(out))
    findings = check_deck.check_slide(1, Presentation(str(out)).slides[0], prs, 40)
    messages = " ".join(f.message for f in findings if f.level == "error")
    assert "ナレーション" in messages and "12pt" in messages


def test_example_builds_clean(tmp_path):
    import build_example

    out = tmp_path / "sample.pptx"
    build_example.build().save(str(out))
    assert check_deck.main([str(out)]) == 0


def test_proposal_tag_is_visible_on_white():
    sl = Deck().slide("見出し", NARRATION)
    tag = sl.tag(1, 1, "提案")
    assert str(tag.fill.fore_color.rgb) == "1F3A5F"
